"""Read a supplied master workbook into a private snapshot, never a web asset.
Usage: python extract-product-master.py SOURCE.xlsx OUTPUT.json
OUTPUT must be outside this repository (which is published via GitHub Pages).
"""
import collections
import datetime
import hashlib
import json
import pathlib
import re
import sys
import openpyxl

source = pathlib.Path(sys.argv[1]).resolve()
target = pathlib.Path(sys.argv[2]).resolve()
repo = pathlib.Path(__file__).resolve().parent.parent
if target.is_relative_to(repo):
    raise SystemExit('Private master output must be outside the published repository')
before = hashlib.sha256(source.read_bytes()).hexdigest()
wb = openpyxl.load_workbook(source, data_only=True)
ws = wb['제품 마스터']
headers = [c.value for c in ws[6]]
assert len(headers) == 52 and all(isinstance(h, str) for h in headers)
assert headers[10] == '마스터 제품 코드' and headers[30] == '바코드'
date_match = re.search(r'자료 기준 (\d{4}-\d{2}-\d{2})', ws['A2'].value)
assert date_match, 'Missing source date'
def value(cell):
    v = cell.value
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d')
    return v
products = []
row_count = 0
for cells in ws.iter_rows(min_row=7):
    if all(c.value is None for c in cells):
        continue
    assert cells[0].value is not None, 'Unexpected nameless source row'
    row = cells[0].row
    entry = {'row': row, 'values': [value(c) for c in cells]}
    if ws.row_dimensions[row].outlineLevel == 0:
        assert cells[11].value in ('완제품', '기획세트', 'GWP')
        assert cells[10].value, 'Missing product identity'
        products.append({**entry, 'components': []})
    else:
        assert ws.row_dimensions[row].outlineLevel == 1 and products
        products[-1]['components'].append(entry)
    row_count += 1
component_count = sum(len(p['components']) for p in products)
assert len(products) == 141 and component_count == 1473 and row_count == 1614
flat = [entry for p in products for entry in (p, *p['components'])]
assert [r['row'] for r in flat] == list(range(7, 1621))
snapshot = {'id': True, 'source_name': source.name, 'source_date': date_match[1],
            'source_sha256': before, 'headers': headers, 'products': products}
target.parent.mkdir(parents=True, exist_ok=True)
target.write_text(json.dumps(snapshot, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
saved = json.loads(target.read_text(encoding='utf-8'))
assert saved == snapshot
assert hashlib.sha256(source.read_bytes()).hexdigest() == before
print(json.dumps({'products': len(products), 'components': component_count, 'rows': row_count,
                  'columns': len(headers), 'source_sha256': before, 'bytes': target.stat().st_size,
                  'source_unchanged': True, 'round_trip_equal': True}, ensure_ascii=False))
